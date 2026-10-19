"""
Export Module for LedgerLens.
Exports transactions to CSV and formatted Microsoft Excel (XLSX).
"""

import io
import csv
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_transactions_to_csv(transactions: List[Dict[str, Any]]) -> str:
    """Generate CSV string for transactions."""
    output = io.StringIO()
    fieldnames = ["date", "party", "item", "quantity", "amount", "currency", "type", "confidence", "flagged", "flag_reason"]
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for row in transactions:
        writer.writerow(row)
    return output.getvalue()

def export_transactions_to_xlsx(transactions: List[Dict[str, Any]]) -> bytes:
    """Generate XLSX binary stream with styling and formatted columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ledger Transactions"

    headers = ["Date", "Party / Customer", "Item / Description", "Qty", "Amount (INR)", "Type", "Confidence", "Flagged", "Review Notes"]
    ws.append(headers)

    # Style header row
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Add transaction rows
    flag_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    
    for r_idx, t in enumerate(transactions, 2):
        flagged_val = bool(t.get("flagged") or float(t.get("confidence", 1.0)) < 0.7)
        row_data = [
            str(t.get("date", "")),
            str(t.get("party", "")),
            str(t.get("item", "")),
            float(t.get("quantity", 1.0)),
            float(t.get("amount", 0.0)),
            str(t.get("type", "")).upper(),
            f"{float(t.get('confidence', 1.0)):.2f}",
            "YES" if flagged_val else "NO",
            str(t.get("flag_reason", "") or "")
        ]
        ws.append(row_data)

        # Highlight flagged row
        if flagged_val:
            for c_idx in range(1, len(row_data) + 1):
                ws.cell(row=r_idx, column=c_idx).fill = flag_fill

        # Borders
        for c_idx in range(1, len(row_data) + 1):
            ws.cell(row=r_idx, column=c_idx).border = thin_border

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
